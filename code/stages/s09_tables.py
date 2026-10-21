"""Stage 09 - manuscript tables, generated from pipeline results.

Each table is written as CSV (machine-readable), and all tables together as tables.xlsx (one sheet per table) and
tables.docx (editable Word tables in Elsevier style: caption above, three horizontal rules - top, below the header
and bottom - no vertical rules, footnotes below). Output: manuscript/tables/.

Run:  python run.py tables
"""
from __future__ import annotations
import sys, json
from pathlib import Path
import numpy as np
import pandas as pd
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from kittur.config import load, Stage


def _f(v, nd=1):
    return '-' if v is None or (isinstance(v, float) and np.isnan(v)) else f'{v:.{nd}f}'


def _ci(d, nd=2):
    return f"{d['median']:.{nd}f} ({d['lo95']:.{nd}f}-{d['hi95']:.{nd}f})"


# --------------------------------------------------------------------------------------------- tables
def _items(C, e):
    j = json.load(open(C.f(C['inputs']['epochs'][e] + '.json')))
    return j['items'] if isinstance(j, dict) else j    # [[id, date, cloud], ...]


def t1_sources(C, R):
    fr = C['corona']['frames']
    ep = {e: _items(C, e) for e in C['inputs']['epochs']}
    sc = lambda e: f"{len(ep[e])} scenes, {min(i[1] for i in ep[e])} to {max(i[1] for i in ep[e])}"
    rows = [
        ('Military correspondence', 'National Archives of India, Abhilekh Patal: Foreign Political proceedings', 'Oct 1824-Jan 1825', '-', 'Siege narrative, positions', 'Archival record'),
        ('Engineers\' descriptions and gazetteer', 'Gazetteer of the Bombay Presidency XXI, Belgaum (Campbell, 1884), quoting Lawe (1825) and Welsh (1828)', '1825-1884', '-', 'Fort plan and dimensions', 'Public domain'),
        ('Topographic map', 'Survey of India one-inch sheets 48 I/09, I/10, I/13, I/14 (Zenodo 10.5281/zenodo.18133392)', 'surveyed 1922-23, publ. 1924', '1:63,360', 'Pre-dam villages, place names', 'CC BY 4.0 (Zenodo record)'),
    ]
    for ent, g in fr.items():
        rows.append(('CORONA ' + g['camera'], f"{ent}, piece {g['piece']}", str(g['date']), 'c. 2-3 m (resampled to 2 m)', 'Fort plan, 1967 stereo pair' if '217' in ent else 'Fort plan, pre-dam landscape', 'Public domain (USGS EarthExplorer)'))
    for site, g in C['sites'].items():
        lab = g['label'].replace('CORONA ', '')
        cam, ent, *date = lab.split(' ', 2)[0], lab.split(' ')[1].rstrip(','), lab.split(', ')[-1]
        rows.append((f'CORONA {cam} ({site.capitalize()})', ent, date[0], 'c. 2-3 m (resampled to 2 m)', 'Transferability test', 'Public domain (USGS EarthExplorer)'))
    for e in ('L1990', 'L2000', 'L2010'):
        rows.append(('Landsat 5 TM / 7 ETM+ Collection 2 Level-2', f'Dry-season median ({e})', sc(e), '30 m', 'Change metrics', 'Public domain (USGS)'))
    for e in ('S2016', 'S2020', 'S2026'):
        rows.append(('Sentinel-2 MSI Level-2A', f'Dry-season median ({e})', sc(e), '10-20 m', 'Change metrics', 'Copernicus Sentinel data (free and open)'))
    rows += [('Sentinel-2 MSI Level-2A', 'S2B_43PDT_20260217_0_L2A', '2026-02-17', '10 m', 'Georeferencing reference, 2026 imagery', 'Copernicus Sentinel data (free and open)'),
             ('Digital surface model', 'Copernicus DEM GLO-30', '2011-2015', '30 m', 'Terrain, viewshed, least-cost paths', 'Copernicus DEM licence'),
             ('Vector data', 'OpenStreetMap (Overpass extract)', '2026-10-20', '-', 'Reservoir, tanks, roads, Belgaum fort outline', 'ODbL 1.0'),
             ('Vector data', 'Natural Earth admin-1 and land', 'v5', '1:10m, 1:50m', 'Location map', 'Public domain')]
    d = pd.DataFrame(rows, columns=['Data', 'Product / identifier', 'Date(s)', 'Resolution', 'Use', 'Licence'])
    return d, 'Data sources.', ['CORONA dates are acquisition dates; scan pieces are USGS 7 µm scans. Scene identifiers of the median composites are listed in the epoch JSON files (raw/landsat, raw/sentinel, sites/*/raw/epochs).']


def t2_georef_v1(C, R):   # superseded by t2_georef (stage 12); kept for the record, not written
    A = json.load(open(R / 'georef_accuracy.json')); co = A['coreg_1967']
    n = lambda k: A.get(k, {})
    rows = [
        ('SoI 48 I/14 etc. (1924)', 'Four neatline corners', 'Sheet corners', 4, '-', '-', 'Errors of the order of 100 m (datum and sheet distortion)'),
        ('DA224, Kittur crop', 'Local affine (chained patch matching to Sentinel-2)', 'Independent checkpoints', n('kittur_recorded_checkpoints')['n'], _f(n('kittur_recorded_checkpoints')['rms_m']), '-', 'Valid near the fort; 50-100 m at the crop margins'),
        ('DA224, Kittur crop', '', 'Automatic patch NCC vs Sentinel-2', n('corona_1968_kittur_auto_ncc')['n_kept'], _f(n('corona_1968_kittur_auto_ncc')['rms_m']), f"{np.hypot(n('corona_1968_kittur_auto_ncc')['mean_dE'], n('corona_1968_kittur_auto_ncc')['mean_dN']):.0f}", 'Farmland matches unstable across 60 years'),
        ('DA224, Nandagad crop', 'Local affine', 'Recorded checkpoints', n('nandagad_recorded_checkpoints')['n'], _f(n('nandagad_recorded_checkpoints')['rms_m']), f"{np.hypot(n('nandagad_recorded_checkpoints')['mean_dE'], n('nandagad_recorded_checkpoints')['mean_dN']):.0f}", ''),
        ('DA224, Nandagad crop', '', 'Automatic patch NCC vs Sentinel-2', n('corona_1968_nandagad_auto_ncc')['n_kept'], _f(n('corona_1968_nandagad_auto_ncc')['rms_m']), f"{np.hypot(n('corona_1968_nandagad_auto_ncc')['mean_dE'], n('corona_1968_nandagad_auto_ncc')['mean_dN']):.0f}", ''),
        ('DF224, Sangolli crop', 'Affine transferred from DA224 overlap', 'Automatic patch NCC vs Sentinel-2', n('corona_1968_sangolli_auto_ncc')['n_kept'], _f(n('corona_1968_sangolli_auto_ncc')['rms_m']), f"{np.hypot(n('corona_1968_sangolli_auto_ncc')['mean_dE'], n('corona_1968_sangolli_auto_ncc')['mean_dN']):.0f}", ''),
    ]
    for k, lab in (('DA217', 'DA217 (1967), Kittur'), ('DF217', 'DF217 (1967), Kittur'), ('DF224', 'DF224 (1968), Kittur')):
        g = co[k]
        rows.append((lab, 'SIFT + RANSAC, 2nd-order polynomial', 'Fit residuals relative to the DA224 crop', g['gcps'], _f(g['rms_m']), '-', f"Relative co-registration; coverage {g['coverage']:.0%}"))
    for s in ('sampgaon', 'bailhongal', 'oldsangoli'):
        g = n(f'corona_1968DA223_{s}_auto_ncc')
        rows.append((f'DA223, {s.replace("oldsangoli", "Old Sangoli").capitalize() if s != "oldsangoli" else "Old Sangoli"}', 'Single coarse frame match', 'Automatic patch NCC vs Sentinel-2', g['n_kept'], _f(g['rms_m']), f"{np.hypot(g['mean_dE'], g['mean_dN']):.0f}", 'Presence/absence only'))
    rows.append(('DA223, Tigadi', 'Single coarse frame match', 'Automatic patch NCC vs Sentinel-2', n('corona_1968DA223_tigadi_auto_ncc')['n_patches'], '-', '-', 'Too few matching patches'))
    for site, nm in (('srirangapatna', 'DA028 (1967), Srirangapatna'), ('hampi', 'DA077 (1965), Hampi'), ('belgaum', 'DA222 (1968), Belgaum')):
        f_ = n(f'corona_{site}_recorded_fit'); a_ = n(f'corona_{site}_auto_ncc')
        rows.append((nm, f"Patch-NCC ties + RANSAC affine ({f_['inliers']} of {f_['ties']} ties)", 'Leave-one-out residuals of the ties', f_['inliers'], _f(f_['rms_loo_m']), '-', f"Fit RMS {f_['rms_fit_m']:.1f} m" + ('; provisional' if site == 'hampi' else '')))
        rows.append((nm, '', 'Automatic patch NCC vs Sentinel-2', a_['n_kept'], _f(a_['rms_m']), f"{np.hypot(a_['mean_dE'], a_['mean_dN']):.0f}", ''))
    d = pd.DataFrame(rows, columns=['Dataset', 'Method', 'Check', 'n', 'RMS (m)', 'Mean offset (m)', 'Note'])
    return d, 'Georeferencing and co-registration accuracy.', [
        'Automatic check: 500 m patches at 10 m, normalised cross-correlation (NCC) > 0.45 against Sentinel-2 (17 February 2026 at Kittur), offsets beyond 3 median absolute deviations removed. Its RMS includes matching noise over changed farmland as well as georeferencing error.',
        'Checkpoints were not used in the fits.']


def t3_dimensions_stage07(C, R):   # superseded by t3_dimensions (stage 13); kept for the regression record
    M = json.load(open(R / 'montecarlo.json')); L = json.load(open(R / 'landscape.json'))
    A = M['areas_ha_or_length_m']; rows = []
    for k in ('lower_fort', 'ring_outer', 'ring_middle', 'ring_core', 'pond_nw', 'pond_e'):
        rows.append((A[k]['label'], 'Area (ha)', f"{M['nominal'][k]:.2f}", _ci(A[k])))
    for k in ('moat_nw', 'outer_east'):
        rows.append((A[k]['label'], 'Length (m)', f"{M['nominal'][k]:.0f}", _ci(A[k], 0)))
    rows.append(('Lower fort to citadel core', 'Centroid distance (m)', '-', _ci(M['fort_citadel_distance_m'], 0)))
    H = M['heights_m']
    rows.append(('Citadel mound', 'Relief, Copernicus GLO-30 (m)', '-', _ci(H['Copernicus GLO-30']['citadel'], 1)))
    rows.append(('Citadel mound', 'Relief, 1967 stereo surface (m)', '-', _ci(H['CORONA 1967 stereo']['citadel'], 1)))
    rows.append(('Lower fort', 'Relief, Copernicus GLO-30 (m)', '-', _ci(H['Copernicus GLO-30']['lower'], 1)))
    rows.append(('Lower fort', 'Relief, 1967 stereo surface (m)', '-', _ci(H['CORONA 1967 stereo']['lower'], 1)))
    rows.append(('Citadel summit', 'Visible area within 5 km (km²)', f"{L['citadel']['visible_km2_5km']:.1f}", '-'))
    rows.append(('Lower-fort centre', 'Visible area within 5 km (km²)', f"{L['lower_fort']['visible_km2_5km']:.1f}", '-'))
    d = pd.DataFrame(rows, columns=['Feature', 'Measure', 'As digitised', 'Monte Carlo median (95% interval)'])
    m = M['model']
    return d, 'Dimensions of the Kittur fort features digitised on CORONA DA224 (31 January 1968).', [
        f"Monte Carlo: N = {M['N']}, seed {M['seed']}; {m['translation_rms_m']} m RMS translation, linear term {m['edge_rms_range_m'][0]}-{m['edge_rms_range_m'][1]} m at {m['edge_dist_m'] / 1000:g} km, {m['local_residual_m']} m local residual, {m['vertex_noise_m']} m vertex noise; vertical noise {m['vertical_sd_m']['copernicus']} m (Copernicus) and {m['vertical_sd_m']['stereo']} m (stereo).",
        'Relief: citadel = maximum within 30 m of the core minus the median on 300-600 m circles; lower fort = interior median minus the median on a 450 m circle.',
        'Viewsheds from Copernicus GLO-30 (a surface model), observer 10 m above ground, refraction coefficient 0.13.']


def _commanded(L):
    z0 = L['citadel']['ground_z']
    hi = next(a for a in L['annulus_maxz'] if a['max_z'] > z0 + 1)
    return (f"Copernicus: {L['citadel']['higher_cells_within_1000m']} cells within 1 km exceed the summit by "
            f"{L['citadel']['max_z_within_1000m'] - z0:.1f} m; first ground > 1 m higher: {hi['max_z']:.0f} m at {hi['r0'] / 1000:g}-{hi['r1'] / 1000:g} km; "
            f"{L['citadel']['visible_km2_25km']:.0f} km² visible within 25 km")


def t4_archival(C, R):
    P = yaml.safe_load(open(C.code('params/archival.yaml')))
    M = json.load(open(R / 'montecarlo.json'))['areas_ha_or_length_m']; L = json.load(open(R / 'landscape.json'))
    B = json.load(open(R / 'belgaum_fort.json')); S = json.load(open(R / 'dem_stereo.json'))
    import geopandas as gpd
    from kittur import features
    F = features.build(C); lf = F.loc['lower_fort'].geometry
    dd = [lf.exterior.distance(F.loc['moat_nw'].geometry.interpolate(t, normalized=True)) for t in np.linspace(0, 1, 201)]
    ee = [lf.exterior.distance(F.loc['outer_east'].geometry.interpolate(t, normalized=True)) for t in np.linspace(0, 1, 201)]
    bases = {'Gazetteer': B['gazetteer_ha'], 'OSM': B['osm_way_ha'], '1968 trace': B['inner_wall_line_ha']}
    rng = lambda fr: f"{min(bases.values()) * fr:.1f}-{max(bases.values()) * fr:.1f} ha"
    meas = dict(
        two_part='Two enclosures: lower fort and a separate three-ring citadel mound 530 m to the SW',
        lower_area=f"Implied {rng(P['lawe_fraction_lower'])}; measured lower-fort enclosure {_ci(M['lower_fort'])} ha",
        citadel_area=f"Implied {rng(P['lawe_fraction_citadel'])}; measured outer ring {_ci(M['ring_outer'])} ha, middle ring {_ci(M['ring_middle'])} ha",
        quarry=f"No central depression resolved: stereo core centre minus rim {S['core_centre_minus_rim_m']:+.1f} m (noise c. {S['noise_m']:.0f} m)",
        commanded=_commanded(L),
        ditches=f"80 yd = 73 m. Spacing of the N/W moat line from the enclosure: median {np.median(dd):.0f} m (5-95%: {np.percentile(dd, 5):.0f}-{np.percentile(dd, 95):.0f} m); east outer line {np.median(ee):.0f} m ({np.percentile(ee, 5):.0f}-{np.percentile(ee, 95):.0f} m)",
        advanced_post=f"Not locatable; illustrative zone within 150 m of an assumed East Gate: {L['zones_ha']['Advanced post (150 m of gate)']:.1f} ha",
        east_gate=f"Glacis crest zone 30-100 m from the east face: {L['zones_ha']['Crest of glacis (30-100 m)']:.1f} ha (construction)",
        escalade=f"Gap between the lower fort and the citadel mound: {L['zones_ha']['Escalade gap lower/upper fort']:.1f} ha (construction)",
    )
    rows = [(it['source'], f"\u201c{it['statement']}\u201d", meas[it['measure']]) for it in P['items']]
    rows.append(('Gazetteer 1884 p. 534 (Belgaum fort)', '\u201cabout 1000 yards long and 800 yards broad ... 2900 yards perimeter ... about 100 acres\u201d',
                 f"914 x 732 m, 2652 m; 100 acres = 40.5 ha, but an oval of the stated dimensions encloses {3.14159265/4*914.4*731.5/1e4:.1f} ha. Traced on DA222 (1968): {B['ew_m']:.0f} x {B['ns_m']:.0f} m, moat centreline perimeter {B['perimeter_centreline_m']:.0f} m, inner wall line {B['inner_wall_line_ha']:.1f} ha, moat outer edge {B['moat_outer_edge_ha']:.1f} ha (OSM outline {B['osm_way_ha']:.1f} ha)"))
    d = pd.DataFrame(rows, columns=['Source', 'Statement', 'Measurement or test'])
    return d, 'Archival descriptions of Kittur and Belgaum forts compared with the image measurements.', [
        'Quotations follow the 1884 Gazetteer text and the NAI viewer transcription (OCR, not visually proofread). Implied areas apply Lawe\'s fractions to Belgaum fort taken as 40.5 ha (Gazetteer, 100 acres), 51.8 ha (OpenStreetMap way 138292670) or 58.0 ha (inner wall line traced on CORONA DA222). Zones from Trewman\'s letter are geometric constructions, not located positions. [?] marks an uncertain reading in the transcription.']


def t5_regional(C, R):
    d = pd.read_csv(C.f(C['inputs']['regional_forts']))
    d = d[['site', 'gazetteer', 'frames', 'in_1968', 'in_2026', 'status']]
    d.columns = ['Site', 'Gazetteer reference', 'CORONA frames', '1968', '2026', 'Status']
    return d, 'Forts of the Kittur country named in the 1884 Gazetteer: evidence in 1968 and 2026.', [
        'DA223 placements are coarse (positional uncertainty about 100-200 m) and support presence/absence statements only.']


def t6_temporal(C, R):
    m = pd.read_csv(R / 'epoch_metrics.csv')
    piv = lambda aoi, col: m[m.aoi == aoi].set_index('epoch')[col]
    eps = list(C['inputs']['epochs'])
    meta = {e: _items(C, e) for e in eps}
    d = pd.DataFrame({
        'Epoch': eps, 'Year': [int(m[m.epoch == e].year.iloc[0]) for e in eps],
        'Scenes': [len(meta[e]) for e in eps],
        'Water within 600 m of Old Sangoli (ha)': [_f(piv('oldSangoli_600m', 'water_ha')[e], 1) for e in eps],
        'Water within 3 km (ha)': [_f(piv('sangolli_valley_3km', 'water_ha')[e], 1) for e in eps],
        'NDVI contrast, citadel mound': [_f(piv('ring_outer_contrast', 'ndvi')[e], 3) for e in eps],
        'NDVI contrast, lower fort': [_f(piv('fort_enclosure_contrast', 'ndvi')[e], 3) for e in eps],
        'Bare/built fraction, Kittur 1.5 km': [_f(piv('kittur_1.5km', 'built_frac')[e], 2) for e in eps],
    })
    d.loc[-1] = ['CORONA 1968', 1968, 1, 'river channel only', '-', '-', '-', '-']; d = d.sort_index().reset_index(drop=True)
    return d, 'Dry-season change metrics at Kittur and Old Sangoli, 1968-2026.', [
        'January-March median composites; water MNDWI > 0.1; bare/built NDBI > -0.05 and NDVI < 0.25 (confuses bare soil with built surfaces); NDVI contrast = feature minus an 80 m outer buffer.']


def t7_cross(C, R):
    S = json.load(open(R / 'site_metrics.json')); A = json.load(open(R / 'georef_accuracy.json'))
    E = pd.read_csv(R / 'sites_epoch_metrics.csv'); K = pd.read_csv(R / 'epoch_metrics.csv')
    B = json.load(open(R / 'belgaum_fort.json'))
    def trend(site):
        q = K[K.aoi == 'kittur_1.5km'] if site == 'kittur' else E[(E.site == site) & (E.aoi == 'r1500m')]
        q = q.sort_values('year'); a, b = q.iloc[0], q.iloc[-1]
        return f"{a.built_frac:.2f} -> {b.built_frac:.2f}", f"{a.ndvi:.2f} -> {b.ndvi:.2f}"
    rows = []
    info = {'kittur': ('Kittur', 'DA224 (31 Jan 1968); DA217/DF217 (11 Nov 1967)', 'Chained local affines', 'c. 25-30 m near the fort (checkpoints)', 'Yes (DA217/DF217)',
                       'Lower fort 13.4 ha and three-ring citadel mound; Old Sangoli submerged'),
            'srirangapatna': ('Srirangapatna', 'DA028 (2 Apr 1967)', None, None, 'No', 'Fort at W end of the Kaveri island visible in 1967'),
            'hampi': ('Hampi', 'DA077 (10 Oct 1965)', None, None, 'No (DF071 not located)', 'Granite ridges and post-monsoon river; provisional placement'),
            'belgaum': ('Belgaum', 'DA222 (31 Jan 1968)', None, None, 'No (DF223 not located)', f"Moat ring; inner wall line {B['inner_wall_line_ha']:.1f} ha")}
    for site, (nm, fr, geo_, rms, st, key) in info.items():
        if geo_ is None:
            f_ = A[f'corona_{site}_recorded_fit']
            geo_ = f"Patch-NCC + RANSAC affine ({f_['inliers']}/{f_['ties']} ties)"
            rms = f"{f_['rms_loo_m']:.0f} (leave-one-out)"
        M = S[site]; bt, nt = trend(site)
        rows.append((nm, fr, geo_, rms, st, f"{M['built_2026_pct']:.0f} / {M['veg_2026_pct']:.0f} / {M['water_2026_pct']:.1f}",
                     f"{M['water_1960s_pct']:.1f} / {M['darkveg_1960s_pct']:.0f}", bt, nt, key))
    d = pd.DataFrame(rows, columns=['Site', 'CORONA frame(s)', 'Georeferencing', 'RMS (m)', 'Stereo DEM', '2026 built/bare / vegetation / water (%)',
                                    '1960s dark-tone water / dark vegetation (%)', 'Bare/built fraction 1990 -> 2026 (1.5 km)', 'NDVI 1990 -> 2026 (1.5 km)', 'Main observation'])
    return d, 'Cross-site summary of the workflow.', [
        '2026 classes in the 5 x 5 km crop (Kittur: 2026 dry-season median; others: single scene; see Fig. 9). 1960s classes are panchromatic tone thresholds and are crude. At Hampi bare granite falls in the built/bare class.',
        'Epoch metrics within 1.5 km of each fort centre (Fig. 10).']





SHORT = lambda lab: (lab.replace('CORONA KH-4A ', 'CORONA ').replace('Survey of India ', 'SoI ')
                     .replace('DS1045-1102', '').replace('DS1040-1039', '').replace('DS1025-1071', ''))


def t2_georef(C, R):
    """Table 2 (v3): accuracy of every source before and after stage 12, from results/validation.json only."""
    V = json.load(open(R / 'validation.json')); rows = []
    for name, t in V['targets'].items():
        if t.get('status') != 'ok':
            continue
        m = t['models'][t['chosen']]; b = m['check_rmse_boot95']
        rows.append((SHORT(t['label']), 'transfer' if C['validation']['targets'][name].get('transfer') else 'Kittur region',
                     t['n_ties'], f"{t['before']['rmse_all_m']:.0f}", t['chosen'], f"{m['block_cv_rmse_m']:.0f}", f"{m['loo_rmse_m']:.0f}",
                     f"{m['check_rmse_m']:.0f} ({b[0]:.0f}-{b[1]:.0f})", f"{m['moran_I_E'][1]:.3f} / {m['moran_I_N'][1]:.3f}"))
    d = pd.DataFrame(rows, columns=['Source', 'Role', 'Ties', 'Before (m)', 'Model', 'Block-CV (m)', 'LOO (m)', 'Check (95% boot.)', "Moran's I p (E/N)"])
    return d, 'Georeferencing accuracy before (recorded transform) and after re-estimation from automatic open-data tie points.', [
        'Before: RMSE of the recorded transform at all ties. Block-CV: leave-spatial-block-out RMSE of the chosen model, the headline figure. LOO: leave-one-out RMSE. Check: RMSE at the held-out ~30% of ties, with a bootstrap 95% interval.',
        f"Model selection rule: {V.get('selection_rule', '')}",
        "Moran's I: two-sided permutation p-values of the LOO residuals; with 18 tests, p < 0.003 is significant after Bonferroni correction."]


def t3_dimensions(C, R):
    """Table 3 (v3): dimensions from stage 13 (empirical error budget), nominal = chosen stage-12 model."""
    P = json.load(open(R / 'propagation.json')); M = P['metrics']; lab = {}
    import yaml as _y
    for f in _y.safe_load(open(C.code('params/features_1968.yaml')))['features']:
        lab[f['id']] = f.get('name', f['id'])
    rows = []
    ci = lambda d, nd: f"{d['v2']:.{nd}f}", lambda d, nd: f"{d['median']:.{nd}f} ({d['lo95']:.{nd}f}-{d['hi95']:.{nd}f})"
    for k, d in M.items():
        if k.startswith('area_ha_') or k.startswith('length_m_'):
            fid = k.split('_', 2)[2]; unit = 'Area (ha)' if k.startswith('area') else 'Length (m)'; nd = 2 if unit == 'Area (ha)' else 0
            rows.append((lab.get(fid, fid), unit, ci[0](d, nd), ci[1](d, nd)))
    for k, a, b_, nd in (('fort_citadel_m', 'Lower fort to citadel core', 'Centroid distance (m)', 0),
                         ('citadel_bearing_deg', 'Lower fort to citadel core', 'Bearing (deg)', 0),
                         ('iou_fort_osm', 'Lower fort vs OSM outline', 'Intersection over union', 2),
                         ('fort_centroid_to_osm_m', 'Lower fort vs OSM outline', 'Centroid offset (m)', 0),
                         ('vis5_lower_km2', 'Lower-fort centre', 'Visible area within 5 km (km2)', 1)):
        rows.append((a, b_, ci[0](M[k], nd), ci[1](M[k], nd)))
    rows.append(('Citadel summit', 'Elevation, GLO-30 (m)', f"{M['summit_z']['v2']:.1f}", 'fixed by the DSM (see text)'))
    rows.append(('Citadel summit', 'Visible area within 5 km (km2)', f"{M['vis5_citadel_km2']['v2']:.1f}", 'fixed by the DSM (see text)'))
    d = pd.DataFrame(rows, columns=['Feature', 'Measure', 'Nominal', 'Median (95% interval)'])
    eb = P['error_budget']; tr = eb['threshold_sensitivity']['target_range_m']
    return d, 'Dimensions of the Kittur features digitised on CORONA DA224 (31 January 1968), with the full error budget.', [
        f"Nominal: chosen stage-12 model ({P['kittur_model']['kind']}, {P['kittur_model']['n_ties']} ties). Interval: {P['n']} realisations, each a bootstrap transform plus local error topped up to a target RMSE drawn from {tr[0]:.0f}-{tr[1]:.0f} m (tie-threshold sensitivity on the block-CV scale), {eb['vertex_noise_m']} m vertex noise and a {eb['digitising_sd_m']:.0f} m systematic line offset per feature.",
        'Viewsheds: 720 rays to 10 km on GLO-30, observer 10 m above ground, refraction coefficient 0.13. The stage-06 map (2880 rays, 25 km) gives slightly larger areas because of finer ray spacing.',
        'The citadel observer is the highest GLO-30 cell inside the outer ring; it is the same cell in every realisation, so its elevation and viewshed carry vertical error only (Section 3.4).']


TABLES = [('Table1_data_sources', t1_sources), ('Table2_georeferencing_accuracy', t2_georef), ('Table3_dimensions', t3_dimensions),
          ('Table4_archival_vs_measured', t4_archival), ('Table5_regional_forts', t5_regional), ('Table6_temporal_metrics', t6_temporal),
          ('Table7_cross_site', t7_cross)]


# --------------------------------------------------------------------------------------------- writers
def _docx_borders(table):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        if edge in ('top', 'bottom'):
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:color'), '000000')
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    tblPr.append(b)
    trPr = table.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)  # repeat header
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = 0; par.paragraph_format.space_before = 0
    for cell in table.rows[0].cells:   # rule below the header
        tcPr = cell._tc.get_or_add_tcPr(); tb = OxmlElement('w:tcBorders'); e = OxmlElement('w:bottom')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '000000'); tb.append(e); tcPr.append(tb)


def write_docx(items, path):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.section import WD_ORIENT
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(10)
    sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, m, Cm(2))
    for k, (name, d, cap, notes) in enumerate(items):
        if k:
            doc.add_page_break()
        p = doc.add_paragraph(); r = p.add_run(f'Table {k + 1}. '); r.bold = True; p.add_run(cap)
        t = doc.add_table(rows=1, cols=len(d.columns)); t.autofit = True
        for j, c in enumerate(d.columns):
            cell = t.rows[0].cells[j]; cell.text = ''; rr = cell.paragraphs[0].add_run(str(c)); rr.bold = True; rr.font.size = Pt(9)
        for _, row in d.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ''; rr = cells[j].paragraphs[0].add_run('' if pd.isna(v) else str(v)); rr.font.size = Pt(9)
        _docx_borders(t)
        for n in notes:
            q = doc.add_paragraph(); q.paragraph_format.space_before = Pt(2); rn = q.add_run(n); rn.font.size = Pt(8)
    doc.save(path)


def write_xlsx(items, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Border, Side, Alignment
    wb = Workbook(); wb.remove(wb.active)
    thin, med = Side(style='thin'), Side(style='medium')
    for k, (name, d, cap, notes) in enumerate(items):
        ws = wb.create_sheet(f'Table{k + 1}')
        ws.cell(1, 1, f'Table {k + 1}. {cap}').font = Font(name='Arial', bold=True, size=9)
        for j, c in enumerate(d.columns, 1):
            x = ws.cell(3, j, c); x.font = Font(name='Arial', bold=True, size=8); x.border = Border(top=med, bottom=thin)
            x.alignment = Alignment(wrap_text=True, vertical='bottom')
        for i, (_, row) in enumerate(d.iterrows(), 4):
            for j, v in enumerate(row, 1):
                x = ws.cell(i, j, None if pd.isna(v) else (v.item() if hasattr(v, 'item') else v)); x.font = Font(name='Arial', size=8)
                x.alignment = Alignment(wrap_text=True, vertical='top')
                if i == 3 + len(d):
                    x.border = Border(bottom=med)
        for j, c in enumerate(d.columns, 1):
            w = max([len(str(c))] + [len(str(v)) for v in d.iloc[:, j - 1]])
            ws.column_dimensions[ws.cell(3, j).column_letter].width = min(max(10, w * 0.9), 60)
        for i, n in enumerate(notes):
            ws.cell(5 + len(d) + i, 1, n).font = Font(name='Arial', size=8, italic=True)
    wb.save(path)


def main(argv=None):
    C = load(); R = C.p('results'); out = C.p('tables'); out.mkdir(parents=True, exist_ok=True)
    with Stage(C, 'tables') as st:
        items = []
        for name, fn in TABLES:
            d, cap, notes = fn(C, R)
            d.to_csv(out / f'{name}.csv', index=False)
            items.append((name, d, cap, notes)); st.note(f'{name}: {len(d)} rows x {len(d.columns)} columns')
        write_docx(items, out / 'tables.docx'); write_xlsx(items, out / 'tables.xlsx')
        (out / 'table_captions.md').write_text('# Table captions and notes\n\n' + '\n'.join(
            f"**Table {k + 1}.** {cap}  \n*File:* `{name}.csv`; sheet `Table{k + 1}` of `tables.xlsx`; `tables.docx` page {k + 1}.  \n" + ''.join(f'Note: {n}  \n' for n in notes)
            for k, (name, d, cap, notes) in enumerate(items)))
        st.note(f'tables.docx, tables.xlsx, {len(items)} CSV -> {out}')
        from stages import _checklist
        ff, nv, xv = _checklist.write(C)
        st.note(f'checklist.md: {len(ff)} figure checks not met {ff if ff else ""}; docx no vertical rules: {nv}; xlsx no vertical borders: {xv}')


if __name__ == '__main__':
    main()
